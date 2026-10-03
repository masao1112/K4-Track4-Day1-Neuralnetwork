"""results_table.py — Lưu kết quả và xuất ra bảng experiments.xlsx."""
from __future__ import annotations

import json
from pathlib import Path
import openpyxl


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result['cfg'], result['history'], result['summary'] ra <results_dir>/<exp_id>.json."""
    Path(results_dir).mkdir(parents=True, exist_ok=True)
    exp_id = result["cfg"]["exp_id"]
    out_path = Path(results_dir) / f"{exp_id}.json"

    data_to_save = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result["summary"],
    }
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data_to_save, f, indent=2, ensure_ascii=False)

    return str(out_path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    p = Path(results_dir)
    if not p.exists():
        return []

    results = []
    for f in sorted(p.glob("*.json")):
        with open(f, "r", encoding="utf-8") as fp:
            results.append(json.load(fp))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary."""
    cfg = result["cfg"]
    summary = result["summary"]
    exp_id = cfg.get("exp_id", "")

    row = {
        "exp_id": exp_id,
        "group": cfg.get("group", ""),
        "description": cfg.get("description", ""),
        "loss": cfg.get("loss", "ce"),
        "optimizer": cfg.get("optimizer", ""),
        "lr": cfg.get("lr", ""),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch", 512),
        "epochs": cfg.get("epochs", 20),
        "hidden": str(cfg.get("hidden", (256, 128))),
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": cfg.get("clip_norm", ""),
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed", 1),
        "step0_loss": summary.get("step0_loss", ""),
        "best_val_loss": summary.get("best_val_loss", ""),
        "best_epoch": summary.get("best_epoch", ""),
        "final_train_loss": summary.get("final_train_loss", ""),
        "final_val_loss": summary.get("final_val_loss", ""),
        "val_acc": summary.get("val_acc", ""),
        "val_macro_f1": summary.get("val_macro_f1", ""),
        "time_per_epoch_s": summary.get("time_per_epoch_s", ""),
        "peak_mem_MB": summary.get("peak_mem_MB", 0.0),
        "diverged": "Có" if summary.get("diverged", False) else "Không",
        "eval_acc": eval_scores.get("eval_acc", "") if eval_scores else "",
        "eval_macro_f1": eval_scores.get("eval_macro_f1", "") if eval_scores else "",
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes,
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet 'Experiments' của mẫu."""
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]

    headers = [cell.value for cell in ws[1]]
    header_to_col = {h: i + 1 for i, h in enumerate(headers) if h is not None}

    # Bắt đầu ghi từ dòng 2
    for idx, r_data in enumerate(rows):
        row_num = idx + 2
        for k, v in r_data.items():
            if k in header_to_col:
                col = header_to_col[k]
                ws.cell(row=row_num, column=col, value=v)

        # Công thức cho 4 cột cuối (AD, AE, AF, AG)
        ws.cell(row=row_num, column=30, value=f'=IF(P{row_num}="","",P{row_num}-LN(7))')
        ws.cell(row=row_num, column=31, value=f'=IF(OR(T{row_num}="",S{row_num}=""),"",T{row_num}-S{row_num})')
        ws.cell(row=row_num, column=32, value=f'=IF(OR(V{row_num}="",Seeds!$C$8=""),"",V{row_num}-Seeds!$C$8)')
        ws.cell(row=row_num, column=33, value=f'=IF(OR(AF{row_num}="",Seeds!$C$10=""),"",IF(ABS(AF{row_num})>Seeds!$C$10,"Có","Không"))')

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Bảng thí nghiệm đã được lưu vào {out_path} ({len(rows)} dòng)")
